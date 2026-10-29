# -*- coding: utf-8 -*-
"""生成一份**完全合成**的 SCT 样例文件。

仓库里不放任何真实 SCT 导出文件：客户名、产品名、物料编码、价格全部是编造的，
只保留解析器需要的列结构与标记行，用于演示和自动化测试。

用法::

    python examples/make_sample.py            # 生成 examples/sample_sct.xlsx
    python examples/make_sample.py out.xlsx   # 指定输出路径
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

SHEET = "项目配置报价"
COVER = "封面"

# A 列标记 / C 列序号 / D 列描述 / E 总数 / F 成交单价 / G 成交总价
ROWS = [
    ("COL_SORTNO", "COL_DESCRIPTION", "COL_SORTNO", "描述", "总数", "成交单价", "成交总价"),
    ("partType.1,partnumber.DEMO-0001", None, 1, "示例存储阵列 A1_站点1", None, 100000.0, 100000.0),
    ("partType.1,partnumber.DEMO-0101", None, None, "机箱：示例机箱(2U,双控,12*3.5英寸)", 2, 20000.0, 40000.0),
    ("partType.1,partnumber.DEMO-0102", None, None, "CPU：2*示例CPU(32C 2.6GHz)", 2, 10000.0, 20000.0),
    ("partType.1,partnumber.DEMO-0103", None, None, "硬盘：8*8TB 7.2K SATA(3.5英寸)", 8, 5000.0, 40000.0),
    ("partType.0,partnumber.DEMO-0201", None, 2, "示例实施服务_站点1", 1, 5000.0, 5000.0),
    ("partType.0,partnumber.DEMO-0202", None, None, "规划设计与实施服务（合成样例）", 1, 5000.0, 5000.0),
    (",productTotalPrice.1", None, None, None, None, None, 100000.0),
    (",serviceTotalPrice.1", None, None, None, None, None, 5000.0),
    (",projectTotalPrice.1", None, None, None, None, None, 105000.0),
]


def make_sample(path):
    """把合成样例写到 ``path``，返回该路径。"""
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = COVER
    cover["B2"] = "报价书名称：合成样例（非真实项目）"
    cover["B3"] = "客户编码：DEMO-CUSTOMER"
    cover["B4"] = "部门：示例部门"
    cover["B5"] = "本文件由 examples/make_sample.py 生成，仅用于演示与测试。"
    for row in range(2, 6):
        cover.cell(row=row, column=2).font = Font(name="等线")

    ws = wb.create_sheet(SHEET)
    for r, row in enumerate(ROWS, 1):
        for c, value in enumerate(row, 1):
            if value is None:
                continue
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="等线")
            if c in (6, 7) and isinstance(value, float):
                cell.number_format = "0.00"
            if c == 4:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws.column_dimensions["D"].width = 46
    ws.column_dimensions["G"].width = 14

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "sample_sct.xlsx")
    print("saved:", make_sample(target))
