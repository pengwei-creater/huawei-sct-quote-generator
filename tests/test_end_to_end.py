# -*- coding: utf-8 -*-
"""端到端测试：合成 SCT → 解析 JSON → 生成报价单 → 校验数值与格式。"""
import importlib.util
import json
import os
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

from quote_generator import build, parse_file  # noqa: E402


def _load_make_sample():
    spec = importlib.util.spec_from_file_location(
        "make_sample", os.path.join(ROOT, "examples", "make_sample.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture()
def sample_sct(tmp_path):
    make_sample = _load_make_sample()
    return make_sample.make_sample(str(tmp_path / "sample_sct.xlsx"))


def test_parse_totals_and_parents(sample_sct):
    data = parse_file(sample_sct)
    assert data is not None
    assert data["totals"] == {"product": 100000.0, "service": 5000.0, "project": 105000.0}
    assert [p["seq"] for p in data["parents"]] == [1, 2]
    assert data["parents"][0]["description"] == "示例存储阵列 A1_站点1"
    assert len(data["parents"][0]["items"]) == 3
    assert data["parents"][1]["items"][0]["part_type"] == "service"


def test_parse_rejects_non_sct(tmp_path):
    xlsx = tmp_path / "not_sct.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hello"
    wb.save(xlsx)
    assert parse_file(str(xlsx)) is None


def test_build_quote(sample_sct, tmp_path):
    with open(os.path.join(ROOT, "examples", "quote_config.example.json"),
              encoding="utf-8") as fh:
        cfg = json.load(fh)
    cfg["output"] = str(tmp_path / "out.xlsx")

    out, total = build(cfg)
    assert total == pytest.approx(105000.0)
    assert os.path.exists(out)

    ws = openpyxl.load_workbook(out)["sheet"]
    assert ws["A1"].value == "序号"
    assert ws["H1"].value == "备注"
    assert ws["A1"].font.name == "等线"
    assert ws["B2"].value == cfg["rows"][0]["name"]
    assert ws["G2"].value == "=F2*D2"
    # 合计行紧跟数据行，公式覆盖数据区
    assert ws.cell(row=3, column=1).value == "合计"
    assert ws.cell(row=3, column=7).value == "=SUM(G2:G2)"
