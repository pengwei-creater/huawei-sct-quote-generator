# -*- coding: utf-8 -*-
"""SCT（Sales Configuration Tool）报价书解析器。

SCT 导出的 ``客户化报价书_内部价格.xlsx`` 含三个工作表：``封面``、``免责声明``、
``项目配置报价``。本模块只读最后一个表，把父产品、BOM 子项与三个合计行抽成字典。

输出结构::

    {
      "file": "xxx.xlsx",
      "path": "D:/sct/xxx.xlsx",
      "totals": {"product": 0.0, "service": 0.0, "project": 0.0},
      "parents": [
        {
          "seq": 1,
          "description": "OceanStor Pacific 9520（订阅）_站点1",
          "qty": 1,
          "unit_price": 0.0,
          "total": 0.0,
          "items": [
            {"description": "OceanStor Pacific 9520机箱(...)",
             "qty": 3, "part_type": "hardware"}
          ]
        }
      ]
    }

``part_type`` 取值 ``hardware``（A 列含 ``partType.1``）、``service``
（A 列含 ``partType.0``）或 ``unknown``。
"""
import glob
import os

import openpyxl

SHEET = "项目配置报价"

__all__ = ["SHEET", "parse_file", "parse_folder"]


def parse_file(path):
    """解析单个 SCT 报价书；不是 SCT 文件时返回 ``None``。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        return None
    ws = wb[SHEET]

    parents = []
    totals = {"product": None, "service": None, "project": None}
    current = None

    for r in range(1, ws.max_row + 1):
        col_a = str(ws.cell(row=r, column=1).value or "")
        col_c = ws.cell(row=r, column=3).value
        col_d = ws.cell(row=r, column=4).value
        col_e = ws.cell(row=r, column=5).value
        col_f = ws.cell(row=r, column=6).value
        col_g = ws.cell(row=r, column=7).value

        desc = "" if col_d is None else str(col_d).strip()

        # 三个合计行
        if "productTotalPrice" in col_a:
            totals["product"] = col_g
            current = None
            continue
        if "serviceTotalPrice" in col_a:
            totals["service"] = col_g
            current = None
            continue
        if "projectTotalPrice" in col_a:
            totals["project"] = col_g
            current = None
            continue

        # 表头 / 模板标记行
        if desc in ("描述", "序号", "") or str(col_c).startswith("COL_SORTNO"):
            continue

        part_type = (
            "hardware"
            if "partType.1" in col_a
            else ("service" if "partType.0" in col_a else "unknown")
        )

        # 父产品：C 列有序号
        if col_c not in (None, ""):
            current = {
                "seq": col_c,
                "description": desc,
                "qty": col_e,
                "unit_price": col_f,
                "total": col_g,
                "items": [],
            }
            # SCT 父产品行的“总数”常为空，用 总价/单价 推台数供参考
            if col_e in (None, "") and col_f and col_g:
                try:
                    current["qty_inferred"] = round(float(col_g) / float(col_f))
                except (TypeError, ValueError, ZeroDivisionError):
                    pass
            parents.append(current)
            continue

        # BOM 子项
        if current is not None and desc:
            current["items"].append(
                {"description": desc, "qty": col_e, "part_type": part_type}
            )

    return {
        "file": os.path.basename(path),
        "path": path,
        "totals": totals,
        "parents": parents,
    }


def parse_folder(folder):
    """解析目录下所有 ``*.xlsx``，跳过 Excel 临时文件与非 SCT 文件。"""
    results = []
    for f in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
        if os.path.basename(f).startswith("~$"):
            continue
        try:
            parsed = parse_file(f)
        except Exception as exc:  # noqa: BLE001 - 单个文件坏了不影响其它文件
            results.append({"file": os.path.basename(f), "error": str(exc)})
            continue
        if parsed is not None:
            results.append(parsed)
    return results
